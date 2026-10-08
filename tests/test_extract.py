"""Text extraction from file bytes."""

from __future__ import annotations

import io

import docx
import pytest

from arango_connector.extract import effective_mime_type, extract_text, html_to_text

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_pdf(text: str) -> bytes:
    """A minimal one-page PDF showing ``text`` in Helvetica, with a correct xref table."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1))
    out.write(b"startxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_pdf_text() -> None:
    result = extract_text(make_pdf("Quarterly numbers"), "application/pdf", "q2.pdf")
    assert result.text is not None and "Quarterly numbers" in result.text
    assert result.reason is None


def test_docx_text_includes_tables() -> None:
    data = make_docx("Hello Ada", "Second line", table=[["k", "v"], ["revenue", "42"]])
    result = extract_text(data, DOCX_MIME, "notes.docx")
    assert result.text == "Hello Ada\nSecond line\nk\tv\nrevenue\t42"


@pytest.mark.parametrize(
    ("data", "mime", "expected"),
    [
        (b"plain text", "text/plain", "plain text"),
        ("café".encode(), "text/plain", "café"),
        ("café".encode("latin-1"), "text/plain", "café"),
        ("﻿bom".encode(), "text/plain", "bom"),
        (b'{"a": 1}', "application/json", '{"a": 1}'),
        (b"a,b\n1,2", "text/csv", "a,b\n1,2"),
        (b"<p>Hi&nbsp;<b>there</b></p><script>x()</script>", "text/html", "Hi there"),
    ],
)
def test_textual_types(data: bytes, mime: str, expected: str) -> None:
    assert extract_text(data, mime).text == expected


def test_octet_stream_falls_back_to_the_extension() -> None:
    assert effective_mime_type("application/octet-stream", "Report.PDF") == "application/pdf"
    assert effective_mime_type("", "data.csv") == "text/csv"
    assert effective_mime_type("image/png; name=x.png", "x.png") == "image/png"
    result = extract_text(make_pdf("via extension"), "application/octet-stream", "r.pdf")
    assert result.text is not None and "via extension" in result.text


def test_unsupported_type() -> None:
    result = extract_text(b"\x89PNG...", "image/png", "photo.png")
    assert result.text is None and result.reason == "unsupported_type"


def test_malformed_file_is_reported_not_raised() -> None:
    result = extract_text(b"not really a pdf", "application/pdf", "bad.pdf")
    assert result.text is None and result.reason == "extraction_failed"


def test_empty_text() -> None:
    assert extract_text(b"   \n ", "text/plain").reason == "no_text"


def test_truncation() -> None:
    result = extract_text(b"abcdefghij", "text/plain", max_chars=4)
    assert result.text == "abcd" and result.truncated


def test_no_limit_when_zero() -> None:
    result = extract_text(b"abcdefghij", "text/plain", max_chars=0)
    assert result.text == "abcdefghij" and not result.truncated


def test_html_to_text_is_shared() -> None:
    assert html_to_text("<div>a</div><div>b</div>") == "a\nb"
