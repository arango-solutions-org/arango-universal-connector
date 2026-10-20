"""Text extraction from file bytes, for sources that ingest files (Gmail attachments today).

Pure apart from the parsing libraries: no I/O. Every failure is reported as a reason
rather than raised, so one bad file never stops a sync.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from html.parser import HTMLParser

logger = logging.getLogger(__name__)

_PDF = {"application/pdf"}
_DOCX = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
_HTML = {"text/html", "application/xhtml+xml"}
# Treated as text and decoded as-is.
_TEXTUAL = {
    "application/json",
    "application/xml",
    "application/csv",
    "application/x-yaml",
    "application/yaml",
    "application/javascript",
}
_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": next(iter(_DOCX)),
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".xml": "application/xml",
    ".html": "text/html",
    ".htm": "text/html",
    ".yaml": "application/yaml",
    ".yml": "application/yaml",
}


@dataclass(frozen=True)
class Extracted:
    text: str | None
    truncated: bool = False
    # Why there is no text: "unsupported_type", "no_text", "encrypted", "extraction_failed".
    reason: str | None = None


def effective_mime_type(mime_type: str, filename: str) -> str:
    """The MIME type to parse as. Mail clients often send ``application/octet-stream``."""
    mime = (mime_type or "").split(";")[0].strip().lower()
    if mime and mime != "application/octet-stream":
        return mime
    lower = filename.lower()
    for extension, guessed in _EXTENSIONS.items():
        if lower.endswith(extension):
            return guessed
    return mime or "application/octet-stream"


def extract_text(data: bytes, mime_type: str, filename: str = "", max_chars: int = 0) -> Extracted:
    """Extract readable text. ``max_chars`` > 0 caps the result and sets ``truncated``."""
    mime = effective_mime_type(mime_type, filename)
    try:
        if mime in _PDF:
            text = _pdf_text(data)
        elif mime in _DOCX:
            text = _docx_text(data)
        elif mime in _HTML:
            text = html_to_text(_decode(data))
        elif mime.startswith("text/") or mime in _TEXTUAL:
            text = _decode(data)
        else:
            return Extracted(None, reason="unsupported_type")
    except _EncryptedError:
        return Extracted(None, reason="encrypted")
    except Exception as exc:
        # Parsers raise all sorts on malformed files; the attachment is still stored.
        logger.debug("text_extraction_failed mime=%s error_type=%s", mime, type(exc).__name__)
        return Extracted(None, reason="extraction_failed")

    text = text.strip()
    if not text:
        return Extracted(None, reason="no_text")
    if max_chars > 0 and len(text) > max_chars:
        return Extracted(text[:max_chars], truncated=True)
    return Extracted(text)


class _EncryptedError(Exception):
    pass


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            break  # only trust utf-16 with a BOM
        return text
    return data.decode("latin-1")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):  # many "encrypted" PDFs only restrict editing
                raise _EncryptedError
        except Exception as exc:
            raise _EncryptedError from exc
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


# -- HTML ---------------------------------------------------------------------------

_WHITESPACE = re.compile(r"\s+")


class _HTMLText(HTMLParser):
    """Collects readable text from HTML: skips script/style, breaks lines at block tags.

    Source whitespace collapses to single spaces, as a browser renders it. ``<br>`` is
    a hard line break; other block tags only start a new line if one isn't started.
    """

    _BLOCKS = frozenset(
        {
            "p",
            "div",
            "li",
            "tr",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "blockquote",
            "hr",
            "ul",
            "ol",
            "table",
        }
    )
    _SKIP = frozenset({"script", "style", "head", "title"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0

    def _soft_break(self) -> None:
        if self._parts and not self._parts[-1].endswith("\n"):
            self._parts.append("\n")

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "br":
            self._parts.append("\n")
        elif tag in self._BLOCKS:
            self._soft_break()

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCKS:
            self._soft_break()

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._parts.append(_WHITESPACE.sub(" ", data))

    def text(self) -> str:
        lines = (line.strip() for line in "".join(self._parts).split("\n"))
        out: list[str] = []
        for line in lines:
            if line or (out and out[-1]):  # at most one blank line in a row
                out.append(line)
        return "\n".join(out).strip()


def html_to_text(markup: str) -> str:
    parser = _HTMLText()
    parser.feed(markup)
    parser.close()
    return parser.text()
